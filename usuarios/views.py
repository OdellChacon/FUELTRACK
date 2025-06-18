from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
import csv
import json
from xml.etree.ElementTree import Element, SubElement, tostring, fromstring
from io import BytesIO
from reportlab.pdfgen import canvas
from .models import Usuario
from .forms import UsuarioForm
from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator
from auditoria.views import registrar_evento_auditoria

try:
    import openpyxl
except ImportError:
    openpyxl = None

@login_required
def listar_usuarios(request):
    usuarios = Usuario.objects.all()
    # Registrar auditoría: visualización de lista
    registrar_evento_auditoria(request, "Ver", "Usuarios", "Visualizó la lista de usuarios")
    return render(request, 'usuarios/lista_usuarios.html', {'usuarios': usuarios})

@login_required
def registrar_usuario(request):
    if request.method == 'POST':
        form = UsuarioForm(request.POST)
        if form.is_valid():
            usuario_obj = form.save(commit=False)
            password = form.cleaned_data.get('password1')
            usuario_obj.set_password(password)
            usuario_obj.save()
            registrar_evento_auditoria(request, "Crear", "Usuarios", f"Usuario creado: {usuario_obj.email}")
            return redirect('listar_usuarios')
    else:
        form = UsuarioForm()
    return render(request, 'usuarios/registrar_usuario.html', {'form': form})

@login_required
def editar_usuario(request, usuario_id):
    usuario = get_object_or_404(Usuario, id=usuario_id)
    campos_auditar = ['email', 'nombre', 'is_active', 'is_staff']
    if request.method == 'POST':
        form = UsuarioForm(request.POST, instance=usuario)
        if form.is_valid():
            old_values = {field: getattr(usuario, field) for field in campos_auditar}
            usuario_obj = form.save(commit=False)
            # No cambiar la contraseña aquí
            usuario_obj.save()
            new_values = {field: getattr(usuario_obj, field) for field in campos_auditar}
            cambios = []
            for field in campos_auditar:
                if old_values[field] != new_values[field]:
                    cambios.append(f"{field}: {old_values[field]} -> {new_values[field]}")
            detalles = f"Usuario editado: {usuario_obj.email}\n" + ("\n".join(cambios) if cambios else "Sin cambios relevantes.")
            registrar_evento_auditoria(request, "Editar", "Usuarios", detalles)
            return redirect('listar_usuarios')
    else:
        form = UsuarioForm(instance=usuario)
    return render(request, 'usuarios/editar_usuario.html', {'form': form, 'usuario': usuario})

@csrf_exempt
def eliminar_usuario(request, usuario_id):
    if request.method == 'POST':
        try:
            usuario = Usuario.objects.get(id=usuario_id)
            email = usuario.email
            usuario.delete()
            # Registrar auditoría: eliminación individual
            registrar_evento_auditoria(request, "Eliminar", "Usuarios", f"Usuario eliminado: {email}")
            return JsonResponse({'success': True, 'usuario_id': usuario_id, 'message': 'Usuario eliminado correctamente'})
        except Usuario.DoesNotExist:
            return JsonResponse({'success': False, 'message': 'Usuario no encontrado.'}, status=404)
    return JsonResponse({'success': False, 'message': 'Método no permitido.'}, status=405)

@csrf_exempt
def eliminar_usuarios(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            ids = data.get('ids', [])
            if not ids:
                return JsonResponse({'success': False, 'message': 'No se proporcionaron IDs para eliminar.'}, status=400)
            usuarios = Usuario.objects.filter(id__in=ids)
            emails = [u.email for u in usuarios]
            usuarios.delete()
            # Registrar auditoría: eliminación múltiple
            registrar_evento_auditoria(request, "Eliminar", "Usuarios", f"Usuarios eliminados: {', '.join(emails)}")
            return JsonResponse({'success': True, 'message': 'Usuarios eliminados correctamente.'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Error al eliminar usuarios: {str(e)}'}, status=500)
    return JsonResponse({'success': False, 'message': 'Método no permitido.'}, status=405)

@csrf_exempt
def importar_usuarios(request):
    if request.method == 'POST' and request.FILES.get('file'):
        file = request.FILES['file']
        filename = file.name.lower()
        print(f"Archivo recibido: {filename}")  # DEPURACIÓN
        creados = 0
        actualizados = 0
        # CSV
        if filename.endswith('.csv'):
            lines = file.read().decode('utf-8').splitlines()
            print(f"Filas leídas (incluyendo encabezado): {len(lines)}")  # DEPURACIÓN
            reader = csv.reader(lines)
            headers = next(reader, None)
            for row in reader:
                print(f"Procesando fila: {row}")  # DEPURACIÓN
                if len(row) < 4:
                    continue
                email = row[0]
                nombre = row[1]
                is_active = row[2].strip().lower() in ['true', '1', 'sí', 'si']
                is_staff = row[3].strip().lower() in ['true', '1', 'sí', 'si']
                usuario, created = Usuario.objects.get_or_create(email=email, defaults={
                    'nombre': nombre,
                    'is_active': is_active,
                    'is_staff': is_staff
                })
                if created:
                    print(f"Usuario creado: {email}")  # DEPURACIÓN
                    creados += 1
                else:
                    updated = False
                    if usuario.nombre != nombre:
                        usuario.nombre = nombre
                        updated = True
                    if usuario.is_active != is_active:
                        usuario.is_active = is_active
                        updated = True
                    if usuario.is_staff != is_staff:
                        usuario.is_staff = is_staff
                        updated = True
                    if updated:
                        usuario.save()
                        actualizados += 1
                        print(f"Usuario actualizado: {email}")  # DEPURACIÓN
            if creados == 0 and actualizados == 0:
                print("No se creó ni actualizó ningún usuario.")  # DEPURACIÓN
            # Al final de cada importación exitosa:
            registrar_evento_auditoria(request, "Importar", "Usuarios", f"Importación de usuarios desde CSV. Creados: {creados}, Actualizados: {actualizados}")
            return JsonResponse({'success': True, 'creados': creados, 'actualizados': actualizados, 'message': f'Usuarios creados: {creados}, actualizados: {actualizados}'})
        # Excel
        elif (filename.endswith('.xlsx') or filename.endswith('.xls')) and openpyxl:
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = rows[0]
            for row in rows[1:]:
                if row is None or len(row) < 4 or not row[0]:
                    continue
                email = str(row[0])
                nombre = str(row[1])
                is_active = str(row[2]).strip().lower() in ['true', '1', 'sí', 'si']
                is_staff = str(row[3]).strip().lower() in ['true', '1', 'sí', 'si']
                usuario, created = Usuario.objects.get_or_create(email=email, defaults={
                    'nombre': nombre,
                    'is_active': is_active,
                    'is_staff': is_staff
                })
                if created:
                    creados += 1
                else:
                    updated = False
                    if usuario.nombre != nombre:
                        usuario.nombre = nombre
                        updated = True
                    if usuario.is_active != is_active:
                        usuario.is_active = is_active
                        updated = True
                    if usuario.is_staff != is_staff:
                        usuario.is_staff = is_staff
                        updated = True
                    if updated:
                        usuario.save()
                        actualizados += 1
            # Al final de cada importación exitosa:
            registrar_evento_auditoria(request, "Importar", "Usuarios", f"Importación de usuarios desde Excel. Creados: {creados}, Actualizados: {actualizados}")
            return JsonResponse({'success': True, 'creados': creados, 'actualizados': actualizados})
        # JSON
        elif filename.endswith('.json'):
            data = json.load(file)
            for u in data:
                email = u.get('email', '')
                nombre = u.get('nombre', '')
                is_active = bool(u.get('is_active', True))
                is_staff = bool(u.get('is_staff', False))
                usuario, created = Usuario.objects.get_or_create(email=email, defaults={
                    'nombre': nombre,
                    'is_active': is_active,
                    'is_staff': is_staff
                })
                if created:
                    creados += 1
                else:
                    updated = False
                    if usuario.nombre != nombre:
                        usuario.nombre = nombre
                        updated = True
                    if usuario.is_active != is_active:
                        usuario.is_active = is_active
                        updated = True
                    if usuario.is_staff != is_staff:
                        usuario.is_staff = is_staff
                        updated = True
                    if updated:
                        usuario.save()
                        actualizados += 1
            # Al final de cada importación exitosa:
            registrar_evento_auditoria(request, "Importar", "Usuarios", f"Importación de usuarios desde JSON. Creados: {creados}, Actualizados: {actualizados}")
            return JsonResponse({'success': True, 'creados': creados, 'actualizados': actualizados})
        # XML
        elif filename.endswith('.xml'):
            xml_data = file.read().decode('utf-8')
            root = fromstring(xml_data)
            for usuario_elem in root.findall('usuario'):
                email = usuario_elem.find('email').text if usuario_elem.find('email') is not None else ''
                nombre = usuario_elem.find('nombre').text if usuario_elem.find('nombre') is not None else ''
                is_active = (usuario_elem.find('is_active').text.lower() in ['true', '1', 'sí', 'si']) if usuario_elem.find('is_active') is not None else True
                is_staff = (usuario_elem.find('is_staff').text.lower() in ['true', '1', 'sí', 'si']) if usuario_elem.find('is_staff') is not None else False
                usuario, created = Usuario.objects.get_or_create(email=email, defaults={
                    'nombre': nombre,
                    'is_active': is_active,
                    'is_staff': is_staff
                })
                if created:
                    creados += 1
                else:
                    updated = False
                    if usuario.nombre != nombre:
                        usuario.nombre = nombre
                        updated = True
                    if usuario.is_active != is_active:
                        usuario.is_active = is_active
                        updated = True
                    if usuario.is_staff != is_staff:
                        usuario.is_staff = is_staff
                        updated = True
                    if updated:
                        usuario.save()
                        actualizados += 1
            # Al final de cada importación exitosa:
            registrar_evento_auditoria(request, "Importar", "Usuarios", f"Importación de usuarios desde XML. Creados: {creados}, Actualizados: {actualizados}")
            return JsonResponse({'success': True, 'creados': creados, 'actualizados': actualizados})
        # SQL (soporta inserts simples)
        elif filename.endswith('.sql'):
            sql_lines = file.read().decode('utf-8').splitlines()
            import re
            for line in sql_lines:
                line = line.strip()
                if line.lower().startswith('insert into') and 'usuarios_usuario' in line.lower():
                    try:
                        values_part = re.search(r'values\s*\((.*)\);?', line, re.IGNORECASE)
                        if not values_part:
                            continue
                        values_str = values_part.group(1)
                        matches = re.findall(r"'(.*?)'|(\d+)", values_str)
                        values = [m[0] if m[0] else m[1] for m in matches]
                        if len(values) >= 4:
                            email = values[0]
                            nombre = values[1]
                            is_active = values[2].lower() in ['true', '1', 'sí', 'si']
                            is_staff = values[3].lower() in ['true', '1', 'sí', 'si']
                            usuario, created = Usuario.objects.get_or_create(email=email, defaults={
                                'nombre': nombre,
                                'is_active': is_active,
                                'is_staff': is_staff
                            })
                            if created:
                                creados += 1
                            else:
                                updated = False
                                if usuario.nombre != nombre:
                                    usuario.nombre = nombre
                                    updated = True
                                if usuario.is_active != is_active:
                                    usuario.is_active = is_active
                                    updated = True
                                if usuario.is_staff != is_staff:
                                    usuario.is_staff = is_staff
                                    updated = True
                                if updated:
                                    usuario.save()
                                    actualizados += 1
                    except Exception:
                        continue
            # Al final de la importación SQL, fuera del for y del try/except
            registrar_evento_auditoria(request, "Importar", "Usuarios", f"Importación de usuarios desde SQL. Creados: {creados}, Actualizados: {actualizados}")
            return JsonResponse({'success': True, 'creados': creados, 'actualizados': actualizados})
        # PDF (no se puede importar desde PDF)
        elif filename.endswith('.pdf'):
            registrar_evento_auditoria(request, "Importar", "Usuarios", "Intento fallido de importar desde PDF")
            return JsonResponse({'success': False, 'message': 'No se puede importar desde PDF'}, status=400)
        else:
            registrar_evento_auditoria(request, "Importar", "Usuarios", "Formato no soportado para importación")
            return JsonResponse({'success': False, 'message': 'Formato no soportado'}, status=400)
    registrar_evento_auditoria(request, "Importar", "Usuarios", "Archivo no válido o sin datos procesados")
    return JsonResponse({'success': False, 'message': 'Archivo no válido o sin datos procesados'}, status=400)

@login_required
def exportar_usuarios(request, formato):
    usuarios = Usuario.objects.all()
    # ...existing code...
    if formato == 'csv':
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="usuarios.csv"'
        writer = csv.writer(response)
        writer.writerow(['email', 'nombre', 'is_active', 'is_staff'])
        for usuario in usuarios:
            writer.writerow([usuario.email, usuario.nombre, usuario.is_active, usuario.is_staff])
        # Registrar auditoría: exportación
        registrar_evento_auditoria(request, "Exportar", "Usuarios", "Exportación de usuarios en formato CSV")
        return response
    elif formato == 'excel':
        if not openpyxl:
            return JsonResponse({'success': False, 'message': 'openpyxl no instalado'}, status=400)
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(['email', 'nombre', 'is_active', 'is_staff'])
        for usuario in usuarios:
            ws.append([usuario.email, usuario.nombre, usuario.is_active, usuario.is_staff])
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="usuarios.xlsx"'
        wb.save(response)
        # Registrar auditoría: exportación
        registrar_evento_auditoria(request, "Exportar", "Usuarios", "Exportación de usuarios en formato Excel")
        return response
    elif formato == 'json':
        response = HttpResponse(content_type='application/json')
        response['Content-Disposition'] = 'attachment; filename="usuarios.json"'
        data = [{'email': u.email, 'nombre': u.nombre, 'is_active': u.is_active, 'is_staff': u.is_staff} for u in usuarios]
        response.write(json.dumps(data, indent=4, ensure_ascii=False))
        # Registrar auditoría: exportación
        registrar_evento_auditoria(request, "Exportar", "Usuarios", "Exportación de usuarios en formato JSON")
        return response
    elif formato == 'xml':
        response = HttpResponse(content_type='application/xml')
        response['Content-Disposition'] = 'attachment; filename="usuarios.xml"'
        root = Element('usuarios')
        for usuario in usuarios:
            usuario_elem = SubElement(root, 'usuario')
            SubElement(usuario_elem, 'email').text = usuario.email
            SubElement(usuario_elem, 'nombre').text = usuario.nombre
            SubElement(usuario_elem, 'is_active').text = str(usuario.is_active)
            SubElement(usuario_elem, 'is_staff').text = str(usuario.is_staff)
        response.write(tostring(root, encoding='unicode'))
        # Registrar auditoría: exportación
        registrar_evento_auditoria(request, "Exportar", "Usuarios", "Exportación de usuarios en formato XML")
        return response
    elif formato == 'pdf':
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="usuarios.pdf"'
        buffer = BytesIO()
        p = canvas.Canvas(buffer)
        p.drawString(100, 800, "Lista de Usuarios")
        y = 780
        p.setFont("Helvetica", 10)
        p.drawString(100, y, "Email")
        p.drawString(250, y, "Nombre")
        p.drawString(400, y, "Activo")
        p.drawString(470, y, "Staff")
        y -= 20
        for usuario in usuarios:
            p.drawString(100, y, str(usuario.email))
            p.drawString(250, y, str(usuario.nombre))
            p.drawString(400, y, "Sí" if usuario.is_active else "No")
            p.drawString(470, y, "Sí" if usuario.is_staff else "No")
            y -= 18
            if y < 40:
                p.showPage()
                y = 800
        p.save()
        buffer.seek(0)
        response.write(buffer.read())
        # Registrar auditoría: exportación
        registrar_evento_auditoria(request, "Exportar", "Usuarios", "Exportación de usuarios en formato PDF")
        return response
    elif formato == 'sql':
        response = HttpResponse(content_type='application/sql')
        response['Content-Disposition'] = 'attachment; filename="usuarios.sql"'
        lines = []
        for usuario in usuarios:
            # Solo los campos principales y en el orden correcto
            lines.append(
                "INSERT INTO usuarios_usuario (email, nombre, is_active, is_staff) VALUES ('{}', '{}', '{}', '{}');".format(
                    usuario.email.replace("'", "''"),
                    usuario.nombre.replace("'", "''"),
                    int(usuario.is_active),
                    int(usuario.is_staff)
                )
            )
        response.write('\n'.join(lines))
        # Registrar auditoría: exportación
        registrar_evento_auditoria(request, "Exportar", "Usuarios", "Exportación de usuarios en formato SQL")
        return response
    registrar_evento_auditoria(request, "Exportar", "Usuarios", f"Intento de exportar en formato no soportado: {formato}")
    return JsonResponse({'success': False, 'message': 'Formato no soportado'}, status=400)

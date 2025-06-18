from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from .models import Estacion
from .forms import EstacionForm
import csv
import json
from xml.etree.ElementTree import Element, SubElement, tostring
from io import BytesIO
from reportlab.pdfgen import canvas
from openpyxl import load_workbook, Workbook
import io
from datetime import datetime, date
from auditoria.models import Auditoria  # Asegúrate de que el modelo se llama así y está importado correctamente
from django.contrib.auth.decorators import login_required
from auditoria.views import registrar_evento_auditoria

def listar_estaciones(request):
    mes = int(request.GET.get('mes', date.today().month))
    anio = int(request.GET.get('anio', date.today().year))
    # Asegúrate de filtrar por mes y año correctamente
    estaciones = Estacion.objects.filter(mes=mes, anio=anio)
    meses = [
        (1, "Enero"), (2, "Febrero"), (3, "Marzo"), (4, "Abril"),
        (5, "Mayo"), (6, "Junio"), (7, "Julio"), (8, "Agosto"),
        (9, "Septiembre"), (10, "Octubre"), (11, "Noviembre"), (12, "Diciembre")
    ]
    anios = list(range(2023, 2031))
    return render(request, 'estaciones/lista_estaciones.html', {
        'estaciones': estaciones,
        'mes': mes,
        'anio': anio,
        'meses': meses,
        'anios': anios,
    })

def registrar_auditoria(request, accion, modulo, detalles):
    # Asigna la instancia de usuario si está autenticado, si no, None
    usuario = request.user if request.user.is_authenticated else None
    Auditoria.objects.create(
        usuario=usuario,
        accion=accion,
        modulo=modulo,
        detalles=detalles
    )

@login_required
def registrar_estacion(request):
    if request.method == 'POST':
        form = EstacionForm(request.POST)
        if form.is_valid():
            estacion = form.save(commit=False)
            estacion.mes = int(request.POST.get('mes', date.today().month))
            estacion.anio = int(request.POST.get('anio', date.today().year))
            estacion.save()
            detalles = "\n".join([f"{getattr(estacion, field)}" for field in form.fields])
            registrar_auditoria(
                request,
                accion='Crear',
                modulo='Estaciones',
                detalles=detalles
            )
            return redirect('listar_estaciones')
    else:
        form = EstacionForm()
    return render(request, 'estaciones/registrar_estacion.html', {'form': form})

@login_required
def editar_estacion(request, estacion_id):
    mes = int(request.GET.get('mes', date.today().month))
    anio = int(request.GET.get('anio', date.today().year))
    estacion = get_object_or_404(Estacion, estacion_id=estacion_id, mes=mes, anio=anio)
    campos_auditar = [
        'region', 'mercado', 'nombre', 'distribuidora', 'tipo_estacion', 'estado',
        'marca', 'modelo', 'serial', 'capacidad', 'consumo', 'pac',
        'tanque_base', 'tanque_externo', 'capacidad_anterior', 'compens', 'fecha'
    ]
    if request.method == 'POST':
        form = EstacionForm(request.POST, instance=estacion)
        if form.is_valid():
            old_values = {field: getattr(estacion, field) for field in campos_auditar}
            estacion = form.save(commit=False)
            estacion.mes = mes
            estacion.anio = anio
            estacion.save()
            new_values = {field: getattr(estacion, field) for field in campos_auditar}
            cambios = []
            for field in campos_auditar:
                if old_values[field] != new_values[field]:
                    cambios.append(f"{field}: {old_values[field]} ----> {new_values[field]}")
            # Cambia aquí el mensaje para que siempre quede claro qué se hizo
            if cambios:
                detalles = f"Estación editada (ID: {estacion.estacion_id}, Nombre: {estacion.nombre}):\n" + "\n".join(cambios)
            else:
                detalles = f"Estación editada (ID: {estacion.estacion_id}, Nombre: {estacion.nombre}): sin cambios en campos principales."
            registrar_auditoria(
                request,
                accion='Editar',
                modulo='Estaciones',
                detalles=detalles
            )
            return redirect('listar_estaciones')
    else:
        form = EstacionForm(instance=estacion)
    return render(request, 'estaciones/editar_estacion.html', {'form': form, 'estacion': estacion})

@login_required
def eliminar_estacion(request, estacion_id):
    estacion = get_object_or_404(Estacion, id=estacion_id)
    detalles = "\n".join([f"{getattr(estacion, field.name)}" for field in Estacion._meta.fields if field.name != "id"])
    estacion.delete()
    registrar_auditoria(
        request,
        accion='Eliminar',
        modulo='Estaciones',
        detalles=detalles
    )
    return JsonResponse({'success': True, 'estacion_id': estacion_id, 'message': 'Estación eliminada correctamente'})

@csrf_exempt
def importar_estaciones(request, formato):
    """
    Importa estaciones desde un archivo en el formato especificado.
    Formatos soportados: excel, csv, json, xml
    """
    if request.method == 'POST' and request.FILES.get('file'):
        file = request.FILES['file']
        errores = []
        from datetime import date
        mes_actual = int(request.GET.get("mes") or request.POST.get("mes") or date.today().month)
        anio_actual = int(request.GET.get("anio") or request.POST.get("anio") or date.today().year)
        try:
            if formato == 'excel':
                workbook = load_workbook(file, data_only=True)
                hoja = workbook.active
                encabezados = [str(cell.value).strip().upper() for cell in next(hoja.iter_rows(min_row=1, max_row=1))]
                mapeo = {
                    "ID": "estacion_id",
                    "REGIÓN": "region",
                    "MERCADO": "mercado",
                    "NOMBRE": "nombre",
                    "DISTRIBUIDORA": "distribuidora",
                    "TIPO DE ESTACIÓN": "tipo_estacion",
                    "ESTADO": "estado",
                    "MARCA": "marca",
                    "MODELO": "modelo",
                    "SERIAL": "serial",
                    "CAPACIDAD2": "capacidad",
                    "COMSUMO": "consumo",
                    "PAC": "pac",
                    "TANQUE BASE": "tanque_base",
                    "TANQUE EXTERNO": "tanque_externo",
                    "TOTAL": "total",
                    "CAPACIDAD ANTERIOR": "capacidad_anterior",
                    "COMPENS": "compens",
                    "NIVEL TOTAL": "nivel_total",
                    "5 MIN": "cinco_min",
                    "10 MIN": "diez_min",
                    "20 MIN": "veinte_min",
                    "FECHA": "fecha"
                }
                procesados = 0
                saltados = 0
                for index, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
                    try:
                        datos = {}
                        for idx, valor in enumerate(fila):
                            encabezado = encabezados[idx]
                            campo = mapeo.get(encabezado)
                            if campo:
                                datos[campo] = valor
                        estacion_id = datos.get('estacion_id')
                        if not estacion_id or str(estacion_id).strip() == "":
                            saltados += 1
                            continue
                        # Conversión de campos numéricos y fechas
                        datos['capacidad'] = parse_float(datos.get('capacidad'))
                        datos['consumo'] = parse_float(datos.get('consumo'))
                        datos['tanque_base'] = parse_float(datos.get('tanque_base'))
                        datos['tanque_externo'] = parse_float(datos.get('tanque_externo'))
                        datos['capacidad_anterior'] = parse_float(datos.get('capacidad_anterior'))
                        datos['compens'] = parse_float(datos.get('compens'))
                        datos['fecha'] = datos.get('fecha') or datetime.now().date()
                        datos['mes'] = mes_actual
                        datos['anio'] = anio_actual
                        Estacion.objects.update_or_create(
                            estacion_id=estacion_id,
                            mes=mes_actual,
                            anio=anio_actual,
                            defaults=datos
                        )
                        print(f"Estacion importada: estacion_id={estacion_id}, mes={mes_actual}, anio={anio_actual}")
                        procesados += 1
                    except Exception as e:
                        errores.append({'fila': index, 'error': str(e)})
                print(f"Estaciones procesadas: {procesados}, saltadas: {saltados}, errores: {len(errores)}")
                registrar_evento_auditoria(request, "Importar", "Estaciones", "Importación de estaciones desde Excel")
            elif formato == 'csv':
                decoded = file.read().decode('utf-8').splitlines()
                reader = csv.reader(decoded)
                headers = next(reader, None)
                for index, fila in enumerate(reader, start=2):
                    try:
                        estacion_id = fila[0]
                        region = fila[1] or "Desconocido"
                        mercado = fila[2] or "Desconocido"
                        nombre = fila[3] or "Sin Nombre"
                        distribuidora = fila[4] or "Desconocido"
                        tipo_estacion = fila[5] or "Desconocido"
                        estado = fila[6] or "Desconocido"
                        marca = fila[7] or "Desconocido"
                        modelo = fila[8] or "Desconocido"
                        capacidad = parse_float(fila[10])
                        consumo = parse_float(fila[11])
                        pac = fila[12] or "Desconocido"
                        tanque_base = parse_float(fila[13])
                        tanque_externo = parse_float(fila[14])
                        capacidad_anterior = parse_float(fila[16])
                        compens = parse_float(fila[17])
                        fecha = fila[22] if len(fila) > 22 and fila[22] else datetime.now().date()
                        Estacion.objects.update_or_create(
                            estacion_id=estacion_id,
                            mes=mes_actual,
                            anio=anio_actual,
                            defaults={
                                'region': region,
                                'mercado': mercado,
                                'nombre': nombre,
                                'distribuidora': distribuidora,
                                'tipo_estacion': tipo_estacion,
                                'estado': estado,
                                'marca': marca,
                                'modelo': modelo,
                                'capacidad': capacidad,
                                'consumo': consumo,
                                'pac': pac,
                                'tanque_base': tanque_base,
                                'tanque_externo': tanque_externo,
                                'capacidad_anterior': capacidad_anterior,
                                'compens': compens,
                                'fecha': fecha,
                                'mes': mes_actual,
                                'anio': anio_actual,
                            }
                        )
                    except Exception as e:
                        errores.append({'fila': index, 'error': str(e)})
                registrar_evento_auditoria(request, "Importar", "Estaciones", "Importación de estaciones desde CSV")
            elif formato == 'json':
                data = json.load(file)
                for index, fila in enumerate(data, start=2):
                    try:
                        estacion_id = fila.get('id') or fila.get('estacion_id')
                        region = fila.get('region', "Desconocido")
                        mercado = fila.get('mercado', "Desconocido")
                        nombre = fila.get('nombre', "Sin Nombre")
                        distribuidora = fila.get('distribuidora', "Desconocido")
                        tipo_estacion = fila.get('tipo_estacion', "Desconocido")
                        estado = fila.get('estado', "Desconocido")
                        marca = fila.get('marca', "Desconocido")
                        modelo = fila.get('modelo', "Desconocido")
                        capacidad = parse_float(fila.get('capacidad'))
                        consumo = parse_float(fila.get('consumo'))
                        pac = fila.get('pac', "Desconocido")
                        tanque_base = parse_float(fila.get('tanque_base'))
                        tanque_externo = parse_float(fila.get('tanque_externo'))
                        capacidad_anterior = parse_float(fila.get('capacidad_anterior'))
                        compens = parse_float(fila.get('compens'))
                        fecha = fila.get('fecha') or datetime.now().date()
                        Estacion.objects.update_or_create(
                            estacion_id=estacion_id,
                            mes=mes_actual,
                            anio=anio_actual,
                            defaults={
                                'region': region,
                                'mercado': mercado,
                                'nombre': nombre,
                                'distribuidora': distribuidora,
                                'tipo_estacion': tipo_estacion,
                                'estado': estado,
                                'marca': marca,
                                'modelo': modelo,
                                'capacidad': capacidad,
                                'consumo': consumo,
                                'pac': pac,
                                'tanque_base': tanque_base,
                                'tanque_externo': tanque_externo,
                                'capacidad_anterior': capacidad_anterior,
                                'compens': compens,
                                'fecha': fecha,
                                'mes': mes_actual,
                                'anio': anio_actual,
                            }
                        )
                    except Exception as e:
                        errores.append({'fila': index, 'error': str(e)})
                registrar_evento_auditoria(request, "Importar", "Estaciones", "Importación de estaciones desde JSON")
            elif formato == 'xml':
                import xml.etree.ElementTree as ET
                tree = ET.parse(file)
                root = tree.getroot()
                for index, estacion_elem in enumerate(root.findall('estacion'), start=2):
                    try:
                        estacion_id = estacion_elem.findtext('id') or estacion_elem.findtext('estacion_id')
                        region = estacion_elem.findtext('region', "Desconocido")
                        mercado = estacion_elem.findtext('mercado', "Desconocido")
                        nombre = estacion_elem.findtext('nombre', "Sin Nombre")
                        distribuidora = estacion_elem.findtext('distribuidora', "Desconocido")
                        tipo_estacion = estacion_elem.findtext('tipo_estacion', "Desconocido")
                        estado = estacion_elem.findtext('estado', "Desconocido")
                        marca = estacion_elem.findtext('marca', "Desconocido")
                        modelo = estacion_elem.findtext('modelo', "Desconocido")
                        capacidad = parse_float(estacion_elem.findtext('capacidad'))
                        consumo = parse_float(estacion_elem.findtext('consumo'))
                        pac = estacion_elem.findtext('pac', "Desconocido")
                        tanque_base = parse_float(estacion_elem.findtext('tanque_base'))
                        tanque_externo = parse_float(estacion_elem.findtext('tanque_externo'))
                        capacidad_anterior = parse_float(estacion_elem.findtext('capacidad_anterior'))
                        compens = parse_float(estacion_elem.findtext('compens'))
                        fecha = estacion_elem.findtext('fecha') or datetime.now().date()
                        Estacion.objects.update_or_create(
                            estacion_id=estacion_id,
                            mes=mes_actual,
                            anio=anio_actual,
                            defaults={
                                'region': region,
                                'mercado': mercado,
                                'nombre': nombre,
                                'distribuidora': distribuidora,
                                'tipo_estacion': tipo_estacion,
                                'estado': estado,
                                'marca': marca,
                                'modelo': modelo,
                                'capacidad': capacidad,
                                'consumo': consumo,
                                'pac': pac,
                                'tanque_base': tanque_base,
                                'tanque_externo': tanque_externo,
                                'capacidad_anterior': capacidad_anterior,
                                'compens': compens,
                                'fecha': fecha,
                                'mes': mes_actual,
                                'anio': anio_actual,
                            }
                        )
                    except Exception as e:
                        errores.append({'fila': index, 'error': str(e)})
                registrar_evento_auditoria(request, "Importar", "Estaciones", "Importación de estaciones desde XML")
            else:
                registrar_evento_auditoria(request, "Importar", "Estaciones", f"Formato no soportado: {formato}")
                return JsonResponse({'success': False, 'message': 'Formato no soportado.'})
            if errores:
                registrar_evento_auditoria(request, "Importar", "Estaciones", f"Errores en importación: {errores}")
                return JsonResponse({'success': False, 'message': 'Algunos registros no se pudieron cargar.', 'errores': errores})
            return JsonResponse({'success': True, 'message': 'Importación exitosa.'})
        except Exception as e:
            registrar_evento_auditoria(request, "Importar", "Estaciones", f"Error al procesar el archivo: {str(e)}")
            return JsonResponse({'success': False, 'message': f'Error al procesar el archivo: {str(e)}'})
    registrar_evento_auditoria(request, "Importar", "Estaciones", "No se recibió ningún archivo.")
    return JsonResponse({'success': False, 'message': 'No se recibió ningún archivo.'})

def parse_float(value):
    """Intenta convertir un valor a float, devuelve 0 si no es posible."""
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0

def exportar_estaciones(request, formato):
    """
    Exporta estaciones en el formato especificado.
    Formatos soportados: excel, csv, json, xml
    """
    estaciones = Estacion.objects.all()
    fields = [field.name for field in Estacion._meta.fields]
    if formato == 'csv':
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="estaciones.csv"'
        writer = csv.writer(response)
        writer.writerow(fields)
        for estacion in estaciones:
            writer.writerow([getattr(estacion, field) for field in fields])
        registrar_evento_auditoria(request, "Exportar", "Estaciones", "Exportación de estaciones en formato CSV")
        return response
    elif formato == 'json':
        response = HttpResponse(content_type='application/json')
        response['Content-Disposition'] = 'attachment; filename="estaciones.json"'
        data = list(estaciones.values())
        response.write(json.dumps(data, indent=4, default=str))
        registrar_evento_auditoria(request, "Exportar", "Estaciones", "Exportación de estaciones en formato JSON")
        return response
    elif formato == 'xml':
        response = HttpResponse(content_type='application/xml')
        response['Content-Disposition'] = 'attachment; filename="estaciones.xml"'
        root = Element('estaciones')
        for estacion in estaciones:
            estacion_elem = SubElement(root, 'estacion')
            for field in fields:
                SubElement(estacion_elem, field).text = str(getattr(estacion, field))
        response.write(tostring(root, encoding='unicode'))
        registrar_evento_auditoria(request, "Exportar", "Estaciones", "Exportación de estaciones en formato XML")
        return response
    elif formato == 'excel':
        wb = Workbook()
        ws = wb.active
        ws.append(fields)
        for estacion in estaciones:
            ws.append([getattr(estacion, field) for field in fields])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="estaciones.xlsx"'
        registrar_evento_auditoria(request, "Exportar", "Estaciones", "Exportación de estaciones en formato Excel")
        return response
    registrar_evento_auditoria(request, "Exportar", "Estaciones", f"Intento de exportar en formato no soportado: {formato}")
    return JsonResponse({'success': False, 'message': 'Formato no soportado'}, status=400)

def importar_excel(request):
    if request.method == 'POST' and request.FILES.get('file'):
        excel_file = request.FILES['file']
        try:
            # Cargar el archivo Excel
            workbook = load_workbook(excel_file, data_only=True)
            
            # Seleccionar la hoja de interés
            hoja_nombre = "Hoja1"  # Cambia esto al nombre de la hoja que necesitas
            if hoja_nombre not in workbook.sheetnames:
                return JsonResponse({'success': False, 'message': f'La hoja "{hoja_nombre}" no existe en el archivo.'})

            hoja = workbook[hoja_nombre]

            # Procesar las filas de la hoja
            for fila in hoja.iter_rows(min_row=2, values_only=True):  # Saltar la fila de encabezados
                # Asumiendo que las columnas están en este orden:
                # ID, Región, Mercado, Nombre, Distribuidora, Tipo Estación, Estado, Marca, Modelo, Serial, Capacidad, Consumo, PAC, Tanque Base, Tanque Externo, Capacidad Anterior, Compens, Fecha
                Estacion.objects.update_or_create(
                    id=fila[0],
                    defaults={
                        'region': fila[1],
                        'mercado': fila[2],
                        'nombre': fila[3],
                        'distribuidora': fila[4],
                        'tipo_estacion': fila[5],
                        'estado': fila[6],
                        'marca': fila[7],
                        'modelo': fila[8],
                        'serial': fila[9],
                        'capacidad': fila[10],
                        'consumo': fila[11],
                        'pac': fila[12],
                        'tanque_base': fila[13],
                        'tanque_externo': fila[14],
                        'capacidad_anterior': fila[15],
                        'compens': fila[16],
                        'fecha': fila[17],
                    }
                )

            return JsonResponse({'success': True, 'message': 'Importación exitosa.'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Error al procesar el archivo: {str(e)}'})

    return JsonResponse({'success': False, 'message': 'No se recibió ningún archivo.'})

def eliminar_estaciones(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            ids = data.get('ids', [])
            if not ids:
                return JsonResponse({'success': False, 'message': 'No se proporcionaron IDs para eliminar.'}, status=400)

            # Eliminar las estaciones con los IDs proporcionados
            Estacion.objects.filter(id__in=ids).delete()
            return JsonResponse({'success': True, 'message': 'Registros eliminados correctamente.'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Error al eliminar registros: {str(e)}'}, status=500)
    return JsonResponse({'success': False, 'message': 'Método no permitido.'}, status=405)

@csrf_exempt
def eliminar_estacion_individual(request, estacion_id):
    if request.method == 'POST':
        try:
            estacion = Estacion.objects.get(id=estacion_id)
            detalles = "\n".join([f"{getattr(estacion, field.name)}" for field in Estacion._meta.fields if field.name != "id"])
            estacion.delete()
            registrar_auditoria(
                request,
                accion='Eliminar',
                modulo='Estaciones',
                detalles=detalles
            )
            return JsonResponse({'success': True, 'estacion_id': estacion_id, 'message': 'Estación eliminada correctamente'})
        except Estacion.DoesNotExist:
            return JsonResponse({'success': False, 'message': 'Estación no encontrada.'}, status=404)
    return JsonResponse({'success': False, 'message': 'Método no permitido.'}, status=405)
